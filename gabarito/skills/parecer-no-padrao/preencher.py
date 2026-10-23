"""Preenche o modelo.docx do escritório com a tabela de um cronologia.md.
Uso: python3 preencher.py cronologia.md modelo.docx saida.docx   (precisa de python-docx: pip install python-docx)"""
import copy, re, sys
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def ler(md):
    linhas = [l.strip() for l in open(md, encoding="utf-8") if l.strip()]
    titulo = next((l.lstrip("# ").strip() for l in linhas if l.startswith("#")), "Cronologia")
    tab = [[c.strip() for c in l.strip("|").split("|")] for l in linhas if l.startswith("|")]
    return titulo, [r for r in tab[2:] if len(r) >= 3]


def sombra(celula, cor):
    pr = celula._tc.get_or_add_tcPr()
    for s in pr.findall(qn("w:shd")):
        pr.remove(s)
    s = OxmlElement("w:shd")
    s.set(qn("w:val"), "clear"); s.set(qn("w:color"), "auto"); s.set(qn("w:fill"), cor)
    pr.append(s)


def trocar(doc, marca, texto):
    for p in doc.paragraphs:
        if marca in p.text:
            for r in p.runs[1:]:
                r.text = ""
            p.runs[0].text = texto
            return p


def main(md, modelo, saida):
    titulo, linhas = ler(md)
    doc = Document(modelo)
    trocar(doc, "{{TITULO}}", titulo.upper())
    tabela = doc.tables[0]
    molde = tabela.rows[1]._tr
    for i, dados in enumerate(linhas):
        tr = copy.deepcopy(molde)
        tabela._tbl.append(tr)
        linha = tabela.rows[-1]
        for celula, texto in zip(linha.cells, dados[:3]):
            celula.paragraphs[0].runs[0].text = texto
            sombra(celula, "F1F1F1" if i % 2 else "FFFFFF")
    tabela._tbl.remove(molde)
    pend = [f"{d[0]}: {d[1]}." for d in linhas if "sem fonte" in d[2].lower()]
    texto = " ".join(pend + ["Conferir datas, prazos e valores nos documentos originais antes de usar."])
    trocar(doc, "{{CONFERENCIA}}", texto)
    doc.save(saida)


if __name__ == "__main__":
    main(*sys.argv[1:4])
